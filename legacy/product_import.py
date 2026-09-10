"""Batch import/export products from Excel (.xlsx) or CSV."""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Iterable

from models import Product
from product_categories import category_for_sku

# Canonical template / export headings (row 1).
TEMPLATE_HEADERS = [
    "SKU",
    "Product Name",
    "Details",
    "Qty",
    "Price",
    "Category",
    "Date Stamp",
]

HEADER_ALIASES = {
    "sku": "SKU",
    "code": "SKU",
    "product sku": "SKU",
    "product code": "SKU",
    "product name": "Product Name",
    "name": "Product Name",
    "item": "Product Name",
    "item name": "Product Name",
    "details": "Details",
    "detail": "Details",
    "serial": "Details",
    "serial number": "Details",
    "s/n": "Details",
    "description": "Details",
    "qty": "Qty",
    "quantity": "Qty",
    "stock": "Qty",
    "in stock": "Qty",
    "price": "Price",
    "unit price": "Price",
    "cost": "Price",
    "category": "Category",
    "cat": "Category",
    "type": "Category",
    "date stamp": "Date Stamp",
    "datestamp": "Date Stamp",
    "import date": "Date Stamp",
    "date": "Date Stamp",
    "count date": "Date Stamp",
    "counted": "Date Stamp",
}


class ImportResult:
    def __init__(self):
        self.added = 0
        self.updated = 0
        self.skipped = 0
        self.skipped_dated = 0
        self.errors: list[str] = []

    @property
    def ok_count(self) -> int:
        return self.added + self.updated


def today_date_stamp(d: datetime | None = None) -> str:
    return (d or datetime.now()).strftime("%Y-%m-%d")


def template_path(base: Path | None = None) -> Path:
    from config import ASSETS_DIR

    root = base or ASSETS_DIR
    return Path(root) / "product_import_template.xlsx"


def write_excel_template(path: Path | str) -> Path:
    """Write a blank Excel template with headings (and example rows)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
    ws.append(["60000", "Dell Latitude Laptop", "i5 16GB 512GB", 2, 699.99, "Laptops", ""])
    ws.append(["92001", "ThinkPad X1", "16GB 512GB", 1, 899.00, "Laptops", ""])
    ws.append(["110001", "iPhone 13", "128GB Blue", 3, 499.00, "Cell Phones", ""])

    widths = {"A": 14, "B": 32, "C": 28, "D": 8, "E": 10, "F": 14, "G": 14}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    wb.save(path)
    return path


def write_csv_template(path: Path | str) -> Path:
    """CSV twin of the Excel template."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(TEMPLATE_HEADERS)
        writer.writerow(["60000", "Dell Latitude Laptop", "i5 16GB 512GB", 2, 699.99, "Laptops", ""])
        writer.writerow(["92001", "ThinkPad X1", "16GB 512GB", 1, 899.00, "Laptops", ""])
        writer.writerow(["110001", "iPhone 13", "128GB Blue", 3, 499.00, "Cell Phones", ""])
    return path


def ensure_templates(assets_dir: Path | None = None) -> tuple[Path, Path]:
    from config import ASSETS_DIR

    root = Path(assets_dir) if assets_dir else ASSETS_DIR
    xlsx = write_excel_template(root / "product_import_template.xlsx")
    csv_path = write_csv_template(root / "product_import_template.csv")
    return xlsx, csv_path


def write_products_csv(path: Path | str, products: list[Product]) -> Path:
    """Export full inventory to CSV (includes Category + Date Stamp)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(TEMPLATE_HEADERS)
        for p in products:
            writer.writerow([
                p.sku or "",
                p.name or "",
                p.serial_number or "",
                p.qty,
                p.price,
                p.category or "",
                getattr(p, "import_date", "") or "",
            ])
    return path


def _normalize_header(raw: str) -> str | None:
    key = (raw or "").strip().lower()
    return HEADER_ALIASES.get(key)


def _map_headers(raw_headers: Iterable[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for i, h in enumerate(raw_headers):
        canon = _normalize_header(str(h) if h is not None else "")
        if canon and canon not in mapping:
            mapping[canon] = i
    return mapping


def _cell(row: list, index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    val = row[index]
    if val is None:
        return ""
    return str(val).strip()


def _parse_price(raw: str) -> float:
    clean = raw.replace("$", "").replace(",", "").strip()
    if not clean:
        return 0.0
    return float(clean)


def _parse_qty(raw: str) -> int:
    clean = raw.replace(",", "").strip()
    if not clean:
        return 0
    return int(float(clean))


def _rows_from_xlsx(path: Path) -> tuple[list[str], list[list]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []
    headers = ["" if c is None else str(c) for c in rows[0]]
    data = [list(r) for r in rows[1:]]
    return headers, data


def _rows_from_csv(path: Path) -> tuple[list[str], list[list]]:
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [], []
    return rows[0], rows[1:]


def read_product_rows(path: Path | str) -> list[dict]:
    """Parse a spreadsheet into product dicts. Raises ValueError on bad format."""
    path = Path(path)
    if not path.exists():
        raise ValueError(f"File not found: {path}")

    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        headers, data = _rows_from_xlsx(path)
    elif suffix in (".csv", ".txt"):
        headers, data = _rows_from_csv(path)
    else:
        raise ValueError("Use an Excel (.xlsx) or CSV file. Google Sheets: File → Download → Excel or CSV.")

    mapping = _map_headers(headers)
    if "Product Name" not in mapping and "SKU" not in mapping:
        raise ValueError(
            "Missing required column: Product Name or SKU.\n"
            f"Expected headings: {', '.join(TEMPLATE_HEADERS)}"
        )

    products: list[dict] = []
    for row_num, row in enumerate(data, start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        name = _cell(row, mapping.get("Product Name"))
        sku = _cell(row, mapping.get("SKU"))
        if not name and not sku:
            continue
        products.append({
            "row": row_num,
            "name": name or "Untitled product",
            "sku": sku,
            "serial_number": _cell(row, mapping.get("Details")),
            "qty": _cell(row, mapping.get("Qty")),
            "price": _cell(row, mapping.get("Price")),
            "category": _cell(row, mapping.get("Category")),
            "import_date": _cell(row, mapping.get("Date Stamp")),
        })
    return products


def import_products_from_file(path: Path | str, *, update_existing_by_sku: bool = True) -> ImportResult:
    """
    Import products into the database.

    Only rows with a blank Date Stamp are imported; those rows are stamped with
    today's date (YYYY-MM-DD). Rows that already have a date are skipped.
    """
    from database import add_product, search_products, update_product

    result = ImportResult()
    try:
        rows = read_product_rows(path)
    except ValueError as e:
        result.errors.append(str(e))
        return result

    if not rows:
        result.errors.append("No product rows found. Fill in the template and try again.")
        return result

    stamp = today_date_stamp()

    for row in rows:
        # Import only when Date Stamp is blank
        if (row.get("import_date") or "").strip():
            result.skipped_dated += 1
            result.skipped += 1
            continue

        try:
            price = _parse_price(row["price"])
            qty = _parse_qty(row["qty"])
            if qty < 0:
                raise ValueError("Qty cannot be negative")
            if price < 0:
                raise ValueError("Price cannot be negative")
        except ValueError as e:
            result.skipped += 1
            result.errors.append(f"Row {row['row']}: {e}")
            continue

        category = (row.get("category") or "").strip()
        if not category:
            category = category_for_sku(row["sku"])

        product = Product(
            id=None,
            name=row["name"],
            serial_number=row["serial_number"],
            sku=row["sku"],
            price=price,
            qty=qty,
            category=category,
            created_at="",
            import_date=stamp,
        )

        existing = None
        if update_existing_by_sku and product.sku:
            matches = [
                p for p in search_products(product.sku)
                if (p.sku or "").strip().lower() == product.sku.lower()
            ]
            if matches:
                existing = matches[0]

        if existing:
            product.id = existing.id
            update_product(product)
            result.updated += 1
        else:
            add_product(product)
            result.added += 1

    return result
