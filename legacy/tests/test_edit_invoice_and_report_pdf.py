"""Tests for invoice update (edit sale) and accounting sales-report PDF."""

from pathlib import Path

import pytest

import config
import database
from models import Invoice, InvoiceItem, Product
from sales_report_pdf import (
    aggregate_monthly_totals,
    build_monthly_totals_pdf,
    build_sales_report_pdf,
)


@pytest.fixture()
def db(tmp_path, monkeypatch):
    db_path = tmp_path / "test.db"
    monkeypatch.setattr(config, "DB_PATH", db_path)
    database.init_db()
    return database


def _make_invoice(**kwargs):
    base = dict(
        id=None,
        invoice_number="INV-0001",
        customer_name="Pat",
        customer_phone="555",
        subtotal=100.0,
        tax_rate=0.13,
        tax_amount=13.0,
        total=113.0,
        payment_method="Cash",
        notes="",
        created_at="2026-09-10 12:00:00",
        items=[],
        discount_type="",
        discount_value=0.0,
        discount_amount=0.0,
        discount_timing="before_tax",
        customer_email="",
    )
    base.update(kwargs)
    return Invoice(**base)


def test_update_invoice_restocks_and_rededucts(db):
    pid = db.add_product(
        Product(
            id=None,
            name="Laptop",
            serial_number="A1",
            sku="92001",
            price=100.0,
            qty=10,
            category="Laptops",
        )
    )
    items = [
        InvoiceItem(
            id=None,
            invoice_id=None,
            product_id=pid,
            description="Laptop",
            serial_number="A1",
            qty=2,
            unit_price=100.0,
            line_total=200.0,
        )
    ]
    inv = _make_invoice(subtotal=200.0, tax_amount=26.0, total=226.0)
    inv_id = db.save_invoice(inv, items)
    products = db.search_products("92001")
    assert products[0].qty == 8

    inv.id = inv_id
    new_items = [
        InvoiceItem(
            id=None,
            invoice_id=None,
            product_id=pid,
            description="Laptop",
            serial_number="A1",
            qty=1,
            unit_price=90.0,
            line_total=90.0,
        )
    ]
    inv.subtotal = 90.0
    inv.tax_amount = 11.7
    inv.total = 101.7
    inv.customer_name = "Pat Edited"
    db.update_invoice(inv, new_items)

    products = db.search_products("92001")
    assert products[0].qty == 9  # restock 2, deduct 1

    loaded = db.get_invoice_by_id(inv_id)
    assert loaded is not None
    assert loaded.customer_name == "Pat Edited"
    assert len(loaded.items) == 1
    assert loaded.items[0].qty == 1
    assert loaded.items[0].unit_price == 90.0


def test_void_invoice_restocks_and_blocks_edit(db):
    pid = db.add_product(
        Product(
            id=None,
            name="Phone",
            serial_number="B2",
            sku="110001",
            price=50.0,
            qty=5,
            category="Cell Phones",
        )
    )
    items = [
        InvoiceItem(
            id=None,
            invoice_id=None,
            product_id=pid,
            description="Phone",
            serial_number="B2",
            qty=2,
            unit_price=50.0,
            line_total=100.0,
        )
    ]
    inv = _make_invoice(
        invoice_number="INV-7777",
        subtotal=100.0,
        tax_amount=13.0,
        total=113.0,
    )
    inv_id = db.save_invoice(inv, items)
    assert db.search_products("110001")[0].qty == 3

    voided = db.void_invoice(inv_id)
    assert int(voided.voided) == 1
    assert voided.voided_at
    assert db.search_products("110001")[0].qty == 5

    loaded = db.get_invoice_by_id(inv_id)
    assert int(loaded.voided) == 1

    with pytest.raises(ValueError, match="already voided"):
        db.void_invoice(inv_id)

    inv.id = inv_id
    with pytest.raises(ValueError, match="voided"):
        db.update_invoice(inv, items)


def test_sales_report_pdf_bytes():
    inv = _make_invoice(invoice_number="INV-0099")
    inv.items = [
        InvoiceItem(
            id=1,
            invoice_id=1,
            product_id=None,
            description="Widget",
            serial_number="",
            qty=1,
            unit_price=10.0,
            line_total=10.0,
        )
    ]
    pdf = build_sales_report_pdf(
        [inv],
        period_label="Today",
        start_date="2026-09-10 00:00:00",
        end_date="2026-09-10 23:59:59",
        settings={"business_name": "Test Shop"},
    )
    assert pdf[:4] == b"%PDF"
    assert len(pdf) > 200


def test_aggregate_monthly_totals_skips_voided_and_sums():
    july_a = _make_invoice(
        invoice_number="INV-0701",
        created_at="2026-07-05 10:00:00",
        subtotal=100.0,
        tax_amount=13.0,
        total=113.0,
    )
    july_b = _make_invoice(
        invoice_number="INV-0702",
        created_at="2026-07-20 11:00:00",
        subtotal=50.0,
        tax_amount=6.5,
        total=56.5,
    )
    aug = _make_invoice(
        invoice_number="INV-0801",
        created_at="2026-08-01 09:00:00",
        subtotal=200.0,
        tax_amount=26.0,
        total=226.0,
    )
    voided = _make_invoice(
        invoice_number="INV-0802",
        created_at="2026-08-15 09:00:00",
        subtotal=999.0,
        tax_amount=99.0,
        total=1098.0,
        voided=1,
    )
    rows = aggregate_monthly_totals([july_a, july_b, aug, voided])
    assert [r["month_key"] for r in rows] == ["2026-07", "2026-08"]
    assert rows[0]["invoices"] == 2
    assert rows[0]["revenue"] == pytest.approx(169.5)
    assert rows[0]["tax"] == pytest.approx(19.5)
    assert rows[1]["invoices"] == 1
    assert rows[1]["revenue"] == pytest.approx(226.0)
    assert rows[1]["tax"] == pytest.approx(26.0)


def test_monthly_totals_pdf_bytes():
    invoices = [
        _make_invoice(created_at="2026-06-01 12:00:00", total=10.0, tax_amount=1.0, subtotal=9.0),
        _make_invoice(
            invoice_number="INV-0002",
            created_at="2026-07-15 12:00:00",
            total=20.0,
            tax_amount=2.0,
            subtotal=18.0,
        ),
    ]
    pdf = build_monthly_totals_pdf(
        invoices,
        period_label="Custom: 2026-06-01 → 2026-07-31",
        start_date="2026-06-01 00:00:00",
        end_date="2026-07-31 23:59:59",
        settings={"business_name": "Test Shop"},
    )
    assert pdf[:4] == b"%PDF"
    assert len(pdf) > 200


def test_parse_user_date_formats():
    from utils import parse_report_date

    assert parse_report_date("2026-09-01") == "2026-09-01 00:00:00"
    assert parse_report_date("09/01/2026", end_of_day=True) == "2026-09-01 23:59:59"
    assert parse_report_date("not-a-date") is None


def test_committed_templates_have_date_stamp():
    root = Path(__file__).resolve().parent.parent / "assets"
    csv_text = (root / "product_import_template.csv").read_text(encoding="utf-8")
    assert "Date Stamp" in csv_text.splitlines()[0]
    assert "Category" in csv_text.splitlines()[0]
    from openpyxl import load_workbook

    headers = list(
        next(load_workbook(root / "product_import_template.xlsx").active.iter_rows(max_row=1, values_only=True))
    )
    assert headers[-2:] == ["Category", "Date Stamp"]
