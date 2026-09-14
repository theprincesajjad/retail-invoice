"""Tests for inventory Excel sync (Status / date sold)."""

from pathlib import Path

import pytest

import config
import database
from inventory_excel_sync import (
    STATUS_IN_STOCK,
    STATUS_SOLD,
    SYNC_HEADERS,
    export_inventory_excel,
    product_status,
    pull_inventory_excel,
    read_sync_rows,
    sync_inventory_excel,
    write_inventory_excel,
)
from models import Invoice, InvoiceItem, Product


@pytest.fixture()
def db(tmp_path, monkeypatch):
    db_path = tmp_path / "test.db"
    monkeypatch.setattr(config, "DB_PATH", db_path)
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    database.init_db()
    return database


def _product(**kwargs):
    base = dict(
        id=None,
        name="Laptop",
        serial_number="i5 16GB",
        sku="92001",
        price=500.0,
        qty=1,
        category="Laptops",
        import_date="2026-09-01",
        sold_date="",
    )
    base.update(kwargs)
    return Product(**base)


def test_sync_headers():
    assert SYNC_HEADERS == [
        "SKU",
        "Status",
        "Product Description",
        "Details",
        "qty",
        "Category",
        "price",
        "date added",
        "date sold",
    ]


def test_write_and_read_roundtrip(tmp_path: Path):
    products = [
        _product(sku="92001", qty=1, sold_date=""),
        _product(id=2, sku="110001", name="Phone", qty=0, sold_date="2026-09-10", category="Cell Phones"),
    ]
    path = write_inventory_excel(tmp_path / "inv.xlsx", products)
    rows = read_sync_rows(path)
    assert len(rows) == 2
    assert rows[0]["status"] == STATUS_IN_STOCK
    assert rows[0]["name"] == "Laptop"
    assert rows[1]["status"] == STATUS_SOLD
    assert rows[1]["sold_date"] == "2026-09-10"


def test_sync_pulls_new_sku_then_exports(db, tmp_path: Path):
    path = tmp_path / "inventory.xlsx"
    write_inventory_excel(
        path,
        [
            _product(sku="92NEW", name="From Excel", qty=2, price=199.0, import_date=""),
        ],
    )
    result = sync_inventory_excel(path)
    assert result.added == 1
    assert result.exported >= 1
    found = [p for p in db.list_all_products() if p.sku == "92NEW"]
    assert len(found) == 1
    assert found[0].name == "From Excel"
    assert found[0].qty == 2
    rows = read_sync_rows(path)
    assert any(r["sku"] == "92NEW" and r["status"] == STATUS_IN_STOCK for r in rows)


def test_sale_sets_sold_date_and_excel(db, tmp_path: Path):
    pid = db.add_product(_product(qty=1, sold_date=""))
    inv = Invoice(
        id=None,
        invoice_number="INV-1001",
        customer_name="Pat",
        customer_phone="",
        subtotal=500.0,
        tax_rate=0.0,
        tax_amount=0.0,
        total=500.0,
        payment_method="Cash",
        notes="",
        created_at="2026-09-14 12:00:00",
        items=[],
    )
    items = [
        InvoiceItem(
            id=None,
            invoice_id=None,
            product_id=pid,
            description="Laptop",
            serial_number="",
            qty=1,
            unit_price=500.0,
            line_total=500.0,
        )
    ]
    db.save_invoice(inv, items)
    product = db.search_products("92001")[0]
    assert product.qty == 0
    assert product.sold_date
    assert product_status(product) == STATUS_SOLD

    path = tmp_path / "inv.xlsx"
    db.save_setting("inventory_excel_path", str(path))
    export_inventory_excel(path)
    rows = read_sync_rows(path)
    sold = next(r for r in rows if r["sku"] == "92001")
    assert sold["status"] == STATUS_SOLD
    assert sold["sold_date"] == product.sold_date


def test_void_clears_sold_date(db):
    pid = db.add_product(_product(qty=1))
    inv = Invoice(
        id=None,
        invoice_number="INV-1002",
        customer_name="Pat",
        customer_phone="",
        subtotal=500.0,
        tax_rate=0.0,
        tax_amount=0.0,
        total=500.0,
        payment_method="Cash",
        notes="",
        created_at="2026-09-14 12:00:00",
        items=[],
    )
    items = [
        InvoiceItem(
            id=None,
            invoice_id=None,
            product_id=pid,
            description="Laptop",
            serial_number="",
            qty=1,
            unit_price=500.0,
            line_total=500.0,
        )
    ]
    db.save_invoice(inv, items)
    assert db.search_products("92001")[0].qty == 0
    db.void_invoice(inv.id)
    restored = db.search_products("92001")[0]
    assert restored.qty == 1
    assert (restored.sold_date or "") == ""


def test_pull_status_sold_forces_qty_zero(db, tmp_path: Path):
    db.add_product(_product(sku="92SOLD", qty=1, sold_date=""))
    path = tmp_path / "inv.xlsx"
    write_inventory_excel(
        path,
        [_product(sku="92SOLD", qty=0, sold_date="2026-08-01")],
    )
    # Manually set status column by rewriting with Sold
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    # header row already has Status; ensure data row Status = Sold
    ws.cell(row=2, column=2, value=STATUS_SOLD)
    ws.cell(row=2, column=5, value=0)
    ws.cell(row=2, column=9, value="2026-08-01")
    wb.save(path)

    pull_inventory_excel(path)
    p = [x for x in db.list_all_products() if x.sku == "92SOLD"][0]
    assert p.qty == 0
    assert p.sold_date == "2026-08-01"
